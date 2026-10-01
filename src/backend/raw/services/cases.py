"""Case lifecycle and every human-in-the-loop action."""

import copy
import hashlib
import io
from datetime import UTC, datetime
from typing import Any

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from raw import config_store
from raw.audit import chain
from raw.domain import EXTRACTION_FIELDS, Role, Status, Vote
from raw.models import (
    AssessmentControl,
    ChangeRequest,
    CommitteeVote,
    Control,
    Document,
    Extraction,
    Override,
    User,
)
from raw.observability import metrics
from raw.services.assessment import case_config, deterministic_critique, latest_assessment, latest_extraction, rescore
from raw.services.gates import committee_blockers
from raw.workflow.machine import TransitionError, committee_round, record_decision, transition

ALLOWED_MIME = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/plain": "txt",
    "text/markdown": "md",
}


class ActionError(ValueError):
    pass


def _reason(session: Session, cr: ChangeRequest, reason: str | None) -> str:
    n = case_config(session, cr)["workflow"]["min_override_reason_chars"]
    if not reason or len(reason.strip()) < n:
        raise ActionError(f"a reason of at least {n} characters is required")
    return reason.strip()


def _require_status(cr: ChangeRequest, *statuses: Status) -> None:
    if Status(cr.status) not in statuses:
        raise TransitionError(f"action not allowed in status {cr.status}")


# ---------- intake (product owner) ----------

def create_case(session: Session, user: User, data: dict[str, Any]) -> ChangeRequest:
    year = datetime.now(UTC).year
    n = (session.execute(select(func.count(ChangeRequest.id))).scalar() or 0) + 1
    cr = ChangeRequest(ref=f"CR-{year}-{n:04d}", title=data["title"], description=data["description"],
                       change_type=data.get("change_type") or "other", business_line=data.get("business_line") or "",
                       launch_date=data.get("launch_date"), status=Status.DRAFT.value, created_by=user.id)
    session.add(cr)
    session.flush()
    chain.record(session, user.username, "case.created", "change_request", cr.id,
                 {"ref": cr.ref, "title": cr.title, "change_type": cr.change_type})
    return cr


def parse_document(filename: str, content: bytes, mime: str) -> str:
    kind = ALLOWED_MIME.get(mime) or {"pdf": "pdf", "docx": "docx", "txt": "txt", "md": "md"}.get(filename.rsplit(".", 1)[-1].lower())
    if kind is None:
        raise ActionError("unsupported file type: use PDF, DOCX, TXT or MD")
    if kind == "pdf":
        from pypdf import PdfReader

        return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(content)).pages)
    if kind == "docx":
        import docx

        return "\n".join(p.text for p in docx.Document(io.BytesIO(content)).paragraphs)
    return content.decode("utf-8", errors="replace")


def add_document(session: Session, cr: ChangeRequest, user: User, filename: str, content: bytes, mime: str,
                 max_bytes: int) -> Document:
    _require_status(cr, Status.DRAFT, Status.INFO_REQUESTED)
    if len(content) > max_bytes:
        raise ActionError(f"file exceeds {max_bytes // (1024 * 1024)} MB")
    text = parse_document(filename, content, mime)
    doc = Document(cr_id=cr.id, filename=filename[:200], mime=mime or "", text=text,
                   sha256=hashlib.sha256(content).hexdigest(), size=len(content), uploaded_by=user.id)
    session.add(doc)
    session.flush()
    chain.record(session, user.username, "document.uploaded", "change_request", cr.id,
                 {"document_id": doc.id, "filename": doc.filename, "sha256": doc.sha256, "size": doc.size})
    return doc


def edit_extraction(session: Session, cr: ChangeRequest, user: User, field: str, value: Any, reason: str) -> Extraction:
    """PO corrects the AI extraction before G0, or an analyst corrects it afterwards (new version, needs rerun)."""
    if field not in EXTRACTION_FIELDS:
        raise ActionError(f"unknown field {field}")
    reason = _reason(session, cr, reason)
    ext = latest_extraction(session, cr.id)
    if ext is None:
        raise ActionError("no extraction yet")
    if user.role == Role.PRODUCT_OWNER:
        _require_status(cr, Status.DRAFT, Status.INFO_REQUESTED)
        target = ext
    else:
        _require_status(cr, Status.ANALYST_REVIEW)
        target = Extraction(cr_id=cr.id, version=ext.version + 1, data=copy.deepcopy(ext.data), flags=ext.flags,
                            provider=ext.provider, prompt_version=ext.prompt_version, confirmed_by=ext.confirmed_by,
                            confirmed_at=ext.confirmed_at)
        session.add(target)
    old = copy.deepcopy(target.data[field])
    data = copy.deepcopy(target.data)
    data[field] = {"value": value, "confidence": 1.0, "source_span": f"human: {reason}"[:160], "origin": "human"}
    target.data = data
    session.flush()
    session.add(Override(cr_id=cr.id, target_type="extraction", target_key=field, old_value=old, new_value=value,
                         reason=reason, user_id=user.id))
    metrics.OVERRIDES.labels("extraction", field).inc()
    chain.record(session, user.username, "override.extraction", "change_request", cr.id,
                 {"field": field, "old": old.get("value"), "new": value, "reason": reason, "extraction_version": target.version})
    return target


def confirm_intake(session: Session, cr: ChangeRequest, user: User) -> None:
    """Gate G0."""
    _require_status(cr, Status.DRAFT, Status.INFO_REQUESTED)
    ext = latest_extraction(session, cr.id)
    if ext is None:
        raise ActionError("run extraction before confirming")
    ext.confirmed_by, ext.confirmed_at = user.id, datetime.now(UTC)
    if cr.config_version_id is None:
        cr.config_version_id = config_store.active(session).id
    chain.record(session, user.username, "gate.G0.passed", "change_request", cr.id, {"extraction_version": ext.version})
    transition(session, cr, Status.SUBMITTED, user.username)
    transition(session, cr, Status.IN_ASSESSMENT, "pipeline")
    cr.info_request = None


# ---------- analyst ----------

def _factor(session: Session, cr: ChangeRequest, dim: str):
    a = latest_assessment(session, cr.id)
    if a is None:
        raise ActionError("no assessment")
    f = next((f for f in a.factors if f.dimension == dim), None)
    if f is None:
        raise ActionError(f"unknown dimension {dim}")
    return a, f


def accept_factor(session: Session, cr: ChangeRequest, user: User, dim: str) -> None:
    _require_status(cr, Status.ANALYST_REVIEW)
    a, f = _factor(session, cr, dim)
    f.confirmed, f.confirmed_by = True, user.id
    metrics.AI_FACTOR_DECISIONS.labels("accepted").inc()
    rescore(session, a, case_config(session, cr))
    chain.record(session, user.username, "gate.G1.factor_accepted", "change_request", cr.id,
                 {"dimension": dim, "likelihood": f.likelihood, "impact": f.impact, "origin": f.origin})


def _tier_consequence(session: Session, cr: ChangeRequest, before: str, after: str) -> bool:
    return before != after and case_config(session, cr)["workflow"]["second_analyst_on_tier_change"]


def override_factor(session: Session, cr: ChangeRequest, user: User, dim: str, likelihood: int, impact: int,
                    reason: str) -> Override:
    _require_status(cr, Status.ANALYST_REVIEW)
    reason = _reason(session, cr, reason)
    if not (1 <= likelihood <= 5 and 1 <= impact <= 5):
        raise ActionError("likelihood and impact must be 1-5")
    a, f = _factor(session, cr, dim)
    before = a.residual_tier
    old = {"likelihood": f.likelihood, "impact": f.impact, "origin": f.origin, "rule_ids": f.rule_ids}
    was_ai = f.origin in {"ai", "offline"}
    f.likelihood, f.impact, f.origin = likelihood, impact, "human"
    f.rationale, f.confirmed, f.confirmed_by = reason, True, user.id
    rescore(session, a, case_config(session, cr))
    a.narrative_stale = a.narrative is not None
    o = Override(cr_id=cr.id, assessment_id=a.id, target_type="factor", target_key=dim, old_value=old,
                 new_value={"likelihood": likelihood, "impact": impact}, reason=reason, user_id=user.id,
                 tier_before=before, tier_after=a.residual_tier,
                 requires_ack=_tier_consequence(session, cr, before, a.residual_tier))
    session.add(o)
    session.flush()
    metrics.OVERRIDES.labels("factor", dim).inc()
    if was_ai:
        metrics.AI_FACTOR_DECISIONS.labels("overridden").inc()
    chain.record(session, user.username, "override.factor", "change_request", cr.id,
                 {"override_id": o.id, "dimension": dim, "old": old, "new": o.new_value, "reason": reason,
                  "tier_before": before, "tier_after": a.residual_tier, "requires_second_analyst": o.requires_ack})
    return o


def set_controls(session: Session, cr: ChangeRequest, user: User, controls: list[dict[str, Any]], reason: str) -> Override:
    _require_status(cr, Status.ANALYST_REVIEW)
    reason = _reason(session, cr, reason)
    a = latest_assessment(session, cr.id)
    lib = {c.id: c for c in session.execute(select(Control)).scalars()}
    before = a.residual_tier
    old = [{"control_id": c.control_id, "design": c.design, "operating": c.operating} for c in a.controls]
    for c in list(a.controls):
        session.delete(c)
    session.flush()
    for c in controls:
        if c["control_id"] not in lib:
            raise ActionError(f"unknown control {c['control_id']}")
        d, op = float(c.get("design", lib[c["control_id"]].design)), float(c.get("operating", lib[c["control_id"]].operating))
        if not (0 <= d <= 1 and 0 <= op <= 1):
            raise ActionError("effectiveness values must be between 0 and 1")
        session.add(AssessmentControl(assessment_id=a.id, control_id=c["control_id"], dims=lib[c["control_id"]].dims,
                                      design=d, operating=op, origin="human", matched_on="analyst"))
    session.flush()
    session.expire(a, ["controls"])
    rescore(session, a, case_config(session, cr))
    a.narrative_stale = a.narrative is not None
    o = Override(cr_id=cr.id, assessment_id=a.id, target_type="control", target_key="controls", old_value=old,
                 new_value=controls, reason=reason, user_id=user.id, tier_before=before, tier_after=a.residual_tier,
                 requires_ack=_tier_consequence(session, cr, before, a.residual_tier))
    session.add(o)
    session.flush()
    metrics.OVERRIDES.labels("control", "ALL").inc()
    chain.record(session, user.username, "override.controls", "change_request", cr.id,
                 {"override_id": o.id, "old": old, "new": controls, "reason": reason, "tier_before": before,
                  "tier_after": a.residual_tier, "requires_second_analyst": o.requires_ack})
    return o


def override_narrative(session: Session, cr: ChangeRequest, user: User, dim: str, text: str, reason: str) -> Override:
    _require_status(cr, Status.ANALYST_REVIEW)
    reason = _reason(session, cr, reason)
    a = latest_assessment(session, cr.id)
    narrative = copy.deepcopy(a.narrative or {"summary": "", "sections": []})
    if dim == "SUMMARY":
        old = narrative.get("summary", "")
        narrative["summary"] = text
    else:
        sec = next((x for x in narrative["sections"] if x["dimension"] == dim), None)
        if sec is None:
            sec = {"dimension": dim, "text": ""}
            narrative["sections"].append(sec)
        old = sec["text"]
        sec["text"] = text
    a.narrative = narrative
    a.narrative_stale = False
    a.narrative_signed_by = None
    claims = deterministic_critique(session, a, narrative)
    a.critique = {"claims": claims, "unsupported": sum(c["status"] != "ok" for c in claims),
                  "checked_by": "deterministic (after human edit)"}
    o = Override(cr_id=cr.id, assessment_id=a.id, target_type="narrative", target_key=dim, old_value=old,
                 new_value=text, reason=reason, user_id=user.id)
    session.add(o)
    session.flush()
    metrics.OVERRIDES.labels("narrative", dim).inc()
    chain.record(session, user.username, "override.narrative", "change_request", cr.id,
                 {"override_id": o.id, "dimension": dim, "reason": reason})
    return o


def signoff_narrative(session: Session, cr: ChangeRequest, user: User, warnings_reason: str | None) -> None:
    """Gate G2."""
    _require_status(cr, Status.ANALYST_REVIEW)
    a = latest_assessment(session, cr.id)
    if not a or not a.narrative:
        raise ActionError("no narrative to sign")
    if a.narrative_stale:
        raise ActionError("narrative is stale after an override: redraft or edit it first")
    unsupported = (a.critique or {}).get("unsupported", 0)
    if unsupported:
        warnings_reason = _reason(session, cr, warnings_reason)
    a.narrative_signed_by, a.narrative_signed_at = user.id, datetime.now(UTC)
    a.warnings_ack_reason = warnings_reason
    chain.record(session, user.username, "gate.G2.narrative_signed", "change_request", cr.id,
                 {"assessment_version": a.version, "unsupported_claims": unsupported, "reason": warnings_reason})


def ack_override(session: Session, cr: ChangeRequest, user: User, override_id: int) -> None:
    """Gate G1b: a second analyst confirms a tier-changing override."""
    o = session.get(Override, override_id)
    if o is None or o.cr_id != cr.id or not o.requires_ack:
        raise ActionError("override does not require acknowledgement")
    if o.user_id == user.id:
        raise ActionError("a different analyst must acknowledge this override")
    if o.acked_by:
        raise ActionError("already acknowledged")
    o.acked_by, o.acked_at = user.id, datetime.now(UTC)
    chain.record(session, user.username, "gate.G1b.override_acknowledged", "change_request", cr.id,
                 {"override_id": o.id, "tier_before": o.tier_before, "tier_after": o.tier_after})


def request_info(session: Session, cr: ChangeRequest, user: User, message: str) -> None:
    _require_status(cr, Status.ANALYST_REVIEW)
    if not message or len(message.strip()) < 10:
        raise ActionError("describe the information needed")
    cr.info_request = message.strip()
    transition(session, cr, Status.INFO_REQUESTED, user.username, message=cr.info_request)


def prepare_rerun(session: Session, cr: ChangeRequest, user: User, scope: str) -> None:
    _require_status(cr, Status.ANALYST_REVIEW)
    if scope == "full":
        transition(session, cr, Status.IN_ASSESSMENT, user.username, reason="analyst rerun")
    chain.record(session, user.username, "pipeline.rerun_requested", "change_request", cr.id, {"scope": scope})


def submit_to_committee(session: Session, cr: ChangeRequest, user: User) -> None:
    _require_status(cr, Status.ANALYST_REVIEW)
    blockers = committee_blockers(session, cr)
    if blockers:
        raise ActionError("cannot submit to committee: " + "; ".join(blockers))
    transition(session, cr, Status.COMMITTEE_REVIEW, user.username)


# ---------- committee ----------

def cast_vote(session: Session, cr: ChangeRequest, user: User, vote: str | None, rationale: str, conflict: bool,
              conditions: list[str]) -> CommitteeVote:
    _require_status(cr, Status.COMMITTEE_REVIEW)
    rnd = committee_round(session, cr.id)
    existing = session.execute(select(CommitteeVote).where(
        CommitteeVote.cr_id == cr.id, CommitteeVote.user_id == user.id, CommitteeVote.round == rnd)).scalar_one_or_none()
    if existing:
        raise ActionError("you have already voted in this round")
    if conflict:
        vote, conditions = None, []
    else:
        if vote not in {v.value for v in Vote}:
            raise ActionError("vote must be APPROVE, APPROVE_WITH_CONDITIONS, REJECT or DEFER")
        _reason(session, cr, rationale)
        if vote == Vote.APPROVE_WITH_CONDITIONS and not [c for c in conditions if c.strip()]:
            raise ActionError("list at least one condition")
    v = CommitteeVote(cr_id=cr.id, user_id=user.id, round=rnd, vote=vote, rationale=rationale or "",
                      conflict=conflict, conditions=[c.strip() for c in conditions if c.strip()])
    session.add(v)
    session.flush()
    chain.record(session, user.username, "committee.vote", "change_request", cr.id,
                 {"round": rnd, "vote": vote, "conflict": conflict, "conditions": v.conditions, "rationale": rationale})
    return v


def decide(session: Session, cr: ChangeRequest, user: User):
    return record_decision(session, cr, user, case_config(session, cr)["workflow"])


def return_to_analyst(session: Session, cr: ChangeRequest, user: User, note: str) -> None:
    if user.role != Role.COMMITTEE:
        raise ActionError("only the committee can return a case")
    _reason(session, cr, note)
    transition(session, cr, Status.ANALYST_REVIEW, user.username, note=note)


def reopen_deferred(session: Session, cr: ChangeRequest, user: User) -> None:
    _require_status(cr, Status.DEFERRED)
    transition(session, cr, Status.ANALYST_REVIEW, user.username, reason="reopened after deferral")
